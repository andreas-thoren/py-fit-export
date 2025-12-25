from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table
from py_fit_export.core import FitExporter


def print_key_info(activity):
    extractor = FitExporter(activity)
    print(extractor.extract_key_info())

def export_to_json(activity, out):
    extractor = FitExporter(activity)
    extractor.export_to_json(out)

def append_table_records(activity, workbook, sheet, table):
    wb = load_workbook(workbook)
    ws = wb[sheet]
    tbl: Table = ws.tables[table]
    start_cell, end_cell = tbl.ref.split(":")
    end_col = end_cell.rstrip("0123456789")


if __name__ == "__main__":
    test_activity = Path("ACTIVITY.fit")
    assert test_activity.is_file()

    # append_table_records(test_activity, Path("test_workbook.xlsx"), "Running sessions", "tblRun")
    print_key_info(test_activity)
    # export_to_json(test_activity, Path("activity.json"))
